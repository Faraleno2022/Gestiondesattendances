"""Builds Excel and PDF files from a table described by the browser.

The browser sends text that is already translated, so this module stays language-agnostic.

Table spec:
{
  fileName, title, subtitle?, legend?, generatedAt?, pageLabel? ("Page {page} / {pages}"),
  sheetName?, pageSize? ("A4" | "A3"), fontSize?, freezeColumns?,
  columns: [{ header, width, align?: "left"|"center"|"right", format?: "percent", shade?: bool, highlight?: bool }],
  rows: [[cell]], footer?: [cell]
}
A cell is a string, a number, null, or { value, text?, status? }.
`value` goes into Excel (numbers stay numbers), `text` is what the PDF prints, `status` colours the cell.
"""
import io
import os
import re

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

STATUS_COLORS = {
    "present": ("D9F2E3", "1E8E4E"),
    "absent": ("F9DCD9", "C0392B"),
    "late": ("FBEBD0", "B8730A"),
    "excused": ("DCE8F6", "2F6FB3"),
}
HEADER_FILL = "E9EDF3"
SHADE_FILL = "F0F1F4"
HIGHLIGHT_FILL = "DDE5FA"
BORDER = "C9CFD9"
TEXT = "1D2330"
MUTED = "5D6677"


class SpecError(Exception):
    """The table description sent by the browser is unusable."""


# --- Validation -------------------------------------------------------------

def _str(value, max_length=200):
    return "" if value is None else str(value)[:max_length]


def _number(value, default):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _clamp(value, low, high):
    return min(max(value, low), high)


def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _normalize_cell(cell):
    if cell is None:
        return {"value": None, "text": "", "status": None}
    if isinstance(cell, dict):
        raw = cell.get("value")
        value = raw if raw is None or _is_number(raw) else _str(raw)
        text = cell.get("text")
        status = cell.get("status") if cell.get("status") in STATUS_COLORS else None
        return {"value": value, "text": _str(raw if text is None else text), "status": status}
    if _is_number(cell):
        return {"value": cell, "text": _str(cell), "status": None}
    return {"value": _str(cell), "text": _str(cell), "status": None}


def _normalize_spec(spec):
    if not isinstance(spec, dict):
        raise SpecError()
    columns = spec.get("columns") if isinstance(spec.get("columns"), list) else []
    columns = columns[:80]
    if not columns:
        raise SpecError()
    rows = spec.get("rows") if isinstance(spec.get("rows"), list) else []

    def normalize_row(row):
        row = row if isinstance(row, list) else []
        return [_normalize_cell(row[i] if i < len(row) else None) for i in range(len(columns))]

    def normalize_column(c):
        c = c if isinstance(c, dict) else {}
        return {
            "header": _str(c.get("header"), 100),
            "width": _clamp(_number(c.get("width"), 10), 2, 80),
            "align": c.get("align") if c.get("align") in ("left", "center", "right") else "left",
            "format": "percent" if c.get("format") == "percent" else None,
            "shade": bool(c.get("shade")),
            "highlight": bool(c.get("highlight")),
        }

    return {
        "fileName": re.sub(r'[\\/:*?"<>|]+', "_", _str(spec.get("fileName"), 120)) or "export",
        "title": _str(spec.get("title")),
        "subtitle": _str(spec.get("subtitle")),
        "legend": _str(spec.get("legend"), 400),
        "generatedAt": _str(spec.get("generatedAt")),
        "pageLabel": _str(spec.get("pageLabel")) or "{page} / {pages}",
        "sheetName": re.sub(r"[\\/*?:\[\]]", " ", _str(spec.get("sheetName"), 31)) or "Sheet1",
        "pageSize": "A3" if spec.get("pageSize") == "A3" else "A4",
        "fontSize": _clamp(_number(spec.get("fontSize"), 9), 5, 14),
        "freezeColumns": int(_clamp(_number(spec.get("freezeColumns"), 0), 0, len(columns))),
        "columns": [normalize_column(c) for c in columns],
        "rows": [normalize_row(r) for r in rows[:5000]],
        "footer": normalize_row(spec["footer"]) if isinstance(spec.get("footer"), list) else None,
    }


# --- Excel ------------------------------------------------------------------

def _excel_value(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def build_xlsx(raw_spec):
    spec = _normalize_spec(raw_spec)
    wb = Workbook()
    ws = wb.active
    ws.title = spec["sheetName"]
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if spec["pageSize"] == "A3" else ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    columns = spec["columns"]
    last_col = len(columns)
    thin = Side(style="thin", color=BORDER)
    borders = Border(top=thin, left=thin, bottom=thin, right=thin)

    def fill(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    def write(row_index, col_index, value):
        cell = ws.cell(row=row_index, column=col_index, value=_excel_value(value))
        if isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"  # user text such as a name must never become a formula
        return cell

    # Title block
    row_index = 0
    intro = [
        (spec["title"], Font(bold=True, size=14)),
        (spec["subtitle"], Font(size=11)),
        (spec["legend"], Font(size=10, color=MUTED)),
        (spec["generatedAt"], Font(size=9, italic=True, color=MUTED)),
    ]
    for text, font in intro:
        if not text:
            continue
        row_index += 1
        write(row_index, 1, text).font = font
        ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=last_col)
    if row_index:
        row_index += 1

    row_index += 1
    header_row = row_index
    ws.row_dimensions[header_row].height = 30
    for i, col in enumerate(columns, start=1):
        cell = write(header_row, i, col["header"])
        cell.font = Font(bold=True)
        cell.fill = fill(HIGHLIGHT_FILL if col["highlight"] else HEADER_FILL)
        cell.border = borders
        cell.alignment = Alignment(horizontal=col["align"], vertical="center", wrap_text=True)

    def write_row(cells, bold=False):
        nonlocal row_index
        row_index += 1
        for i, (c, col) in enumerate(zip(cells, columns), start=1):
            cell = write(row_index, i, c["value"])
            cell.border = borders
            cell.alignment = Alignment(horizontal=col["align"], vertical="center")
            if col["format"] == "percent" and _is_number(c["value"]):
                cell.number_format = "0%"
            if c["status"]:
                bg, fg = STATUS_COLORS[c["status"]]
                cell.fill = fill(bg)
                cell.font = Font(bold=True, color=fg)
            elif bold:
                cell.fill = fill(HEADER_FILL)
                cell.font = Font(bold=True)
            elif col["shade"]:
                cell.fill = fill(SHADE_FILL)

    for cells in spec["rows"]:
        write_row(cells)
    if spec["footer"]:
        write_row(spec["footer"], bold=True)

    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = col["width"]
    ws.freeze_panes = ws.cell(row=header_row + 1, column=spec["freezeColumns"] + 1)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), f"{spec['fileName']}.xlsx"


# --- PDF fonts --------------------------------------------------------------

# Fonts covering Latin, accents and Chinese. PDF_FONT / PDF_FONT_BOLD (path to a .ttf/.ttc) override the search.
# Only TrueType outlines work with ReportLab, so OpenType/CFF fonts such as Noto Sans CJK are not listed.
FONT_CANDIDATES = [
    ("C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/msyhbd.ttc"),
    ("C:/Windows/Fonts/simsun.ttc", None),
    ("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc", None),
    ("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", None),
    ("/usr/share/fonts/truetype/arphic/uming.ttc", None),
    ("/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf", None),
    ("/System/Library/Fonts/Supplemental/Arial Unicode.ttf", None),
]


def _register_ttf(name, path):
    try:
        pdfmetrics.registerFont(TTFont(name, path, subfontIndex=0))
        return True
    except Exception:
        return False


def _project_fonts():
    """A .ttf/.ttc file dropped in the project's fonts/ folder takes precedence over system fonts."""
    folder = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")
    if not os.path.isdir(folder):
        return []
    files = sorted(f for f in os.listdir(folder) if f.lower().endswith((".ttf", ".ttc")))
    bold = [f for f in files if "bold" in f.lower() or "bd" in f.lower()]
    regular = [f for f in files if f not in bold]
    if not regular:
        return []
    return [(os.path.join(folder, regular[0]), os.path.join(folder, bold[0]) if bold else None)]


def _load_fonts():
    """Returns (regular, bold, unicode_capable)."""
    env_regular, env_bold = os.environ.get("PDF_FONT"), os.environ.get("PDF_FONT_BOLD")
    candidates = [(env_regular, env_bold)] if env_regular else _project_fonts() + FONT_CANDIDATES
    for regular, bold in candidates:
        if regular and os.path.exists(regular) and _register_ttf("Body", regular):
            has_bold = bool(bold and os.path.exists(bold) and _register_ttf("Body-Bold", bold))
            return "Body", "Body-Bold" if has_bold else "Body", True
    # No system font: Latin text uses Helvetica and Chinese text the built-in Adobe font
    # (STSong-Light), which PDF readers display with their own Chinese font.
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "Helvetica", "Helvetica-Bold", False


FONT_REGULAR, FONT_BOLD, FONT_IS_UNICODE = _load_fonts()


def _font_for(text, bold):
    if FONT_IS_UNICODE:
        return FONT_BOLD if bold else FONT_REGULAR
    try:
        text.encode("cp1252")  # Helvetica covers Western European characters only
        return "Helvetica-Bold" if bold else "Helvetica"
    except UnicodeEncodeError:
        return "STSong-Light"


def _width(text, font, size):
    return pdfmetrics.stringWidth(text, font, size)


def _fit(text, font, size, width):
    """Cuts the text with an ellipsis so it fits on one line."""
    if _width(text, font, size) <= width:
        return text
    while text and _width(text + "…", font, size) > width:
        text = text[:-1]
    return text + "…" if text else ""


_TOKENS = re.compile(r"[\u2e80-\u9fff\uf900-\ufaff\uff00-\uffef]|[^\s\u2e80-\u9fff\uf900-\ufaff\uff00-\uffef]+\s*|\s+")


def _wrap(text, font, size, width):
    """Splits text into lines (on spaces for Latin text, between characters for Chinese)."""
    lines = []
    for paragraph in text.split("\n"):
        line = ""
        for token in _TOKENS.findall(paragraph):
            if _width(line + token, font, size) <= width or not line:
                line += token
            else:
                lines.append(line.rstrip())
                line = token.lstrip()
        lines.append(_fit(line.rstrip(), font, size, width))
    return lines


class _NumberedCanvas(canvas.Canvas):
    """Delays page output so every page can show "page X of N"."""

    def __init__(self, *args, page_label, margin, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_pages = []
        self._page_label = page_label
        self._margin = margin

    def showPage(self):
        self._saved_pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._saved_pages)
        for number, state in enumerate(self._saved_pages, start=1):
            self.__dict__.update(state)
            label = self._page_label.replace("{page}", str(number)).replace("{pages}", str(total))
            font = _font_for(label, False)
            self.setFont(font, 8)
            self.setFillColor(HexColor(f"#{MUTED}"))
            self.drawRightString(self._pagesize[0] - self._margin, self._margin - 14, label)
            super().showPage()
        super().save()


# --- PDF --------------------------------------------------------------------

def build_pdf(raw_spec):
    spec = _normalize_spec(raw_spec)
    page_w, page_h = landscape(A3 if spec["pageSize"] == "A3" else A4)
    margin = 28
    out = io.BytesIO()
    c = _NumberedCanvas(out, pagesize=(page_w, page_h), page_label=spec["pageLabel"], margin=margin)
    c.setTitle(spec["title"])
    c.setCreator("Gestion des présences")

    columns = spec["columns"]
    table_w = page_w - margin * 2
    total = sum(col["width"] for col in columns)
    widths = [col["width"] / total * table_w for col in columns]
    size = spec["fontSize"]
    pad = max(2.0, size * 0.35)
    row_h = size * 1.9
    line_h = size * 1.2
    bottom = margin + 4  # keep clear of the page number

    def rgb(hex_color):
        return HexColor(f"#{hex_color}")

    # Coordinates below are measured from the top of the page, as in the browser.
    def cell(x, y, w, h, lines, align, bold=False, fill=None, color=TEXT):
        if fill:
            c.setFillColor(rgb(fill))
            c.rect(x, page_h - y - h, w, h, stroke=0, fill=1)
        c.setStrokeColor(rgb(BORDER))
        c.setLineWidth(0.4)
        c.rect(x, page_h - y - h, w, h, stroke=1, fill=0)
        lines = [line for line in lines if line] if any(lines) else []
        if not lines:
            return
        c.setFillColor(rgb(color))
        text_h = len(lines) * line_h
        baseline = y + (h - text_h) / 2 + size * 0.95
        for line in lines:
            font = _font_for(line, bold)
            line = _fit(line, font, size, w - pad * 2)
            c.setFont(font, size)
            if align == "center":
                c.drawCentredString(x + w / 2, page_h - baseline, line)
            elif align == "right":
                c.drawRightString(x + w - pad, page_h - baseline, line)
            else:
                c.drawString(x + pad, page_h - baseline, line)
            baseline += line_h

    header_lines = [
        _wrap(col["header"], _font_for(col["header"], True), size, widths[i] - pad * 2)
        for i, col in enumerate(columns)
    ]
    header_h = max(row_h, max(len(lines) for lines in header_lines) * line_h + pad * 2)

    def draw_header(y):
        x = margin
        for i, col in enumerate(columns):
            cell(x, y, widths[i], header_h, header_lines[i], col["align"], bold=True,
                 fill=HIGHLIGHT_FILL if col["highlight"] else HEADER_FILL)
            x += widths[i]
        return y + header_h

    def draw_row(cells, y, bold=False):
        x = margin
        for i, (value, col) in enumerate(zip(cells, columns)):
            status = value["status"]
            if status:
                fill, color = STATUS_COLORS[status]
            else:
                fill = HEADER_FILL if bold else SHADE_FILL if col["shade"] else None
                color = TEXT
            cell(x, y, widths[i], row_h, [value["text"]], col["align"], bold=bold or bool(status), fill=fill, color=color)
            x += widths[i]
        return y + row_h

    # Title block (first page only)
    y = margin
    for text, font_size, bold, color in (
        (spec["title"], 16, True, TEXT),
        (spec["subtitle"], 11, False, TEXT),
        (spec["legend"], 9, False, MUTED),
        (spec["generatedAt"], 8, False, MUTED),
    ):
        if not text:
            continue
        font = _font_for(text, bold)
        c.setFont(font, font_size)
        c.setFillColor(rgb(color))
        y += font_size * 1.25
        c.drawString(margin, page_h - y, _fit(text, font, font_size, table_w))
        y += 3
    y = draw_header(y + 8)

    rows = [(cells, False) for cells in spec["rows"]]
    if spec["footer"]:
        rows.append((spec["footer"], True))
    for cells, bold in rows:
        if y + row_h > page_h - bottom:
            c.showPage()
            y = draw_header(margin)
        y = draw_row(cells, y, bold)

    c.showPage()
    c.save()
    return out.getvalue(), f"{spec['fileName']}.pdf"
